"""보고서 비동기 생성 잡. specs R · 01-상세설계 §3.2.

플로우: 데이터 조립(master+overlay+층별임대) → 가치점수(F-16, 레지스트리 파라미터)
→ python-pptx 생성 → 저장(베타 로컬 / 프로덕션 GCS) → 성공 트랜잭션 안에서 크레딧 차감.
실패 → status=failed + 사유, 크레딧 미차감.

엔트리 2개: run_generate(로컬 BackgroundTasks) · POST /jobs/generate-report(Cloud Tasks).
"""
import json
import os
import datetime as dt
from fastapi import APIRouter
from pydantic import BaseModel
from ..core.db import tx, pool
from ..core.config import settings
from . import value_score, report_calc, use_type

router = APIRouter(prefix="/jobs", tags=["worker"])

REPORT_DIR = os.environ.get("BT_REPORT_DIR", "/tmp/bt-reports")


def _fnum(v) -> float | None:
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


async def _load_formula_params() -> tuple[int, dict[str, float], dict]:
    """반환: (set_version, 수치 파라미터, F-17 시점보정표)."""
    rows = await pool().fetch(
        """SELECT p.set_version, p.formula_id, p.param_key, p.value_num, p.value_json
           FROM ref.formula_params p
           JOIN ref.formula_sets s ON s.set_version = p.set_version AND s.active"""
    )
    params = {r["param_key"]: float(r["value_num"]) for r in rows if r["value_num"] is not None}
    tj = next((r["value_json"] for r in rows if r["param_key"] == "time_adjust"), None)
    time_adjust = json.loads(tj) if isinstance(tj, str) else (tj or {})
    version = rows[0]["set_version"] if rows else 1
    return version, params, time_adjust


async def _assemble(building_pk: str, team_id: int) -> dict:
    """master + 팀 오버레이 병합 + 층별임대 합계."""
    merged = await pool().fetchval("SELECT app.building_view($1,$2)", building_pk, team_id)
    b = json.loads(merged) if isinstance(merged, str) else (merged or {})
    rents = await pool().fetch(
        """SELECT deposit, rent, maintenance, is_vacant FROM app.floor_rents
           WHERE building_pk=$1 AND team_id=$2 AND deleted_at IS NULL""",
        building_pk, team_id,
    )
    b["total_deposit"] = sum(r["deposit"] or 0 for r in rents)
    b["total_rent"] = sum(r["rent"] or 0 for r in rents)
    b["vacant_count"] = sum(1 for r in rents if r["is_vacant"])

    # 가치점수 입력 보강(F-16) — 토지 속성은 대표필지에서, 연수는 날짜→환산
    if b.get("pnu"):
        p = await pool().fetchval(
            """SELECT to_jsonb(x) FROM (
                 SELECT road_frontage, use_zone, shape, slope
                 FROM master.parcels WHERE pnu=$1) x""",
            b["pnu"],
        )
        pj = json.loads(p) if isinstance(p, str) else (p or {})
        for k in ("road_frontage", "use_zone", "shape", "slope"):
            b.setdefault(k, pj.get(k)) if not b.get(k) else None
            if not b.get(k):
                b[k] = pj.get(k)
    if b.get("approval_ymd"):
        try:
            y = dt.date.fromisoformat(str(b["approval_ymd"])[:10])
            b["age_years"] = (dt.date.today() - y).days / 365.25
        except ValueError:
            pass
    if b.get("remodel_ymd"):
        try:
            y = dt.date.fromisoformat(str(b["remodel_ymd"])[:10])
            b["remodel_years"] = (dt.date.today() - y).days / 365.25
        except ValueError:
            pass
    # F-17 수익환원 블렌드용: 마스터 연임대추정 + 자치구 cap rate(구별 2.5~6.9%라 구별 필요)
    inc = await pool().fetchrow(
        """SELECT e.annual_rent,
                  COALESCE(ic.cap, (SELECT cap FROM master.income_cap WHERE gu='_seoul')) AS cap
           FROM master.buildings b
           LEFT JOIN master.building_rent_est e ON e.building_pk = b.building_pk
           LEFT JOIN master.income_cap ic ON ic.gu = substr(b.bjd_code, 1, 5)
           WHERE b.building_pk = $1""",
        building_pk)
    if inc:
        b["est_annual_rent"] = float(inc["annual_rent"]) if inc["annual_rent"] else None
        b["gu_cap"] = float(inc["cap"]) if inc["cap"] else None
    return b


# 공간 필터: market_area 폴리곤 있으면 그 영역, 없으면 center 반경(기본 500m).
_COMP_SPATIAL = """($5::text IS NOT NULL AND ST_Within(b.geom, ST_MakeValid(ST_GeomFromGeoJSON($5::text)))
                    OR $5::text IS NULL AND ST_DWithin(b.geom::geography,
                         ST_SetSRID(ST_MakePoint($1,$2),4326)::geography, $3))"""


def _market_spatial(subject: dict) -> tuple[dict | None, int, dict | None]:
    """market_area 오버레이 → (polygon, radius_m, center). 없으면 (None, 500, None)."""
    ma = subject.get("market_area")
    if isinstance(ma, str):
        try:
            ma = json.loads(ma)
        except json.JSONDecodeError:
            ma = None
    if isinstance(ma, dict):
        if ma.get("kind") == "polygon":
            return ma.get("geojson"), 500, None
        if ma.get("kind") == "circle":
            return None, int(ma.get("radius_m") or 500), ma.get("center")
    return None, 500, None


def _f16_from_ymd(cb: dict) -> dict:
    """approval_ymd/remodel_ymd(date) → age_years/remodel_years 파생(F-16 입력용)."""
    today = dt.date.today()
    for dk, ak in (("approval_ymd", "age_years"), ("remodel_ymd", "remodel_years")):
        if cb.get(dk):
            d = cb[dk] if isinstance(cb[dk], dt.date) else dt.date.fromisoformat(str(cb[dk])[:10])
            cb[ak] = (today - d).days / 365.25
    return cb


def _apply_override(cb: dict, ov: dict) -> dict:
    """S02b 편집값(fields)을 comp 대장값 위에 덮음. 표시값 → F-16 입력 매핑."""
    for k in ("road_frontage", "use_zone", "shape", "slope", "station_dist"):
        if ov.get(k) not in (None, ""):
            cb[k] = ov[k]
    if ov.get("elevator") is not None:
        cb["elevator"] = 1 if ov["elevator"] in (True, 1, "있음") else 0
    if ov.get("approval_ym"):
        ym = str(ov["approval_ym"]).replace("/", "").replace("-", "")[:6]
        if len(ym) >= 6:
            cb["approval_ymd"] = dt.date(int(ym[:4]), int(ym[4:6]), 1)
    if "remodel_ym" in ov:
        rv = str(ov.get("remodel_ym") or "").replace("/", "").replace("-", "")[:6]
        cb["remodel_ymd"] = dt.date(int(rv[:4]), int(rv[4:6]), 1) if len(rv) >= 6 else None
    return cb


# 이용상황(land_use) 섹터 — S01b 필터 taxonomy(이용상황섹터)와 동일. comp 물적 유사성 판정.
_SECTORS = {
    "commercial": {"상업용", "업무용", "상업기타"},      # 상업용 빌딩
    "mixed": {"주상용", "주상기타"},                     # 상가주택
    "resi": {"단독", "연립", "다세대", "아파트", "주거기타"},
    "industrial": {"공업용", "공업기타"},
}
_ADJACENT = {"commercial": {"mixed"}, "mixed": {"commercial"}}   # 소득형 인접(상호 comp 허용, 감점)
_ADJ_FACTOR = 0.7                                                # 인접 섹터 유사도 가중 감점


def _sector_of(lu: str | None) -> str | None:
    return next((s for s, items in _SECTORS.items() if lu in items), None)


def _comp_type_filter(subject_lu: str | None):
    """반환 (allowed_land_uses|None, adjacent_set). None=분류 불가(나지·특수·미분류) → 성격 필터 미적용."""
    sec = _sector_of(subject_lu)
    if sec is None:
        return None, set()
    same = _SECTORS[sec]
    adj = set().union(*[_SECTORS[a] for a in _ADJACENT.get(sec, set())]) if _ADJACENT.get(sec) else set()
    return sorted(same | adj), adj


async def _fetch_comps(building_pk: str, subject: dict, params: dict,
                       overrides: dict | None = None) -> list[dict]:
    """market_area 내 최근 5년 매각사례 + 각 comp F-16 점수. overrides[pk]=편집 fields.
    반환 각 항목에 편집용 fields(대장 원본 표시값) 포함. F-16 결측=0점(게이팅 없음)."""
    overrides = overrides or {}
    poly, radius, center = _market_spatial(subject)
    geom = await pool().fetchrow(
        "SELECT ST_X(geom) AS lng, ST_Y(geom) AS lat FROM master.buildings WHERE building_pk=$1", building_pk)
    clng = (center or {}).get("lng") if center else None
    clat = (center or {}).get("lat") if center else None
    if clng is None and geom:
        clng, clat = geom["lng"], geom["lat"]

    allowed, adj = _comp_type_filter(subject.get("land_use"))   # 성격(섹터) 필터. None=미적용
    rows = await pool().fetch(
        f"""SELECT DISTINCT ON (sh.building_pk)
                  sh.building_pk, sh.contract_ym, sh.price, sh.total_area, sh.land_area,
                  b.gongsi_latest, b.addr, b.land_use,
                  b.road_frontage, b.use_zone, b.shape, b.slope, b.station_dist,
                  b.approval_ymd, b.remodel_ymd, b.elevator,
                  ST_X(b.geom) AS lng, ST_Y(b.geom) AS lat,
                  round(ST_Distance(b.geom::geography,
                        ST_SetSRID(ST_MakePoint($1,$2),4326)::geography)) AS dist_m
           FROM master.sales_history sh
           JOIN master.buildings b ON b.building_pk = sh.building_pk
           WHERE sh.contract_ym >= to_char(now() - interval '5 years', 'YYYYMM')
             AND sh.building_pk <> $4 AND sh.price > 0 AND sh.total_area > 0
             AND ($6::text[] IS NULL OR b.land_use = ANY($6))
             AND {_COMP_SPATIAL}
           ORDER BY sh.building_pk, sh.contract_ym DESC""",
        clng, clat, radius, building_pk, json.dumps(poly) if poly else None, allowed,
    )
    comps = []
    for r in rows:
        cb = dict(r)
        fields = {  # 편집용 표시값(대장 원본)
            "road_frontage": cb["road_frontage"], "station_dist": cb["station_dist"],
            "use_zone": cb["use_zone"], "shape": cb["shape"], "slope": cb["slope"],
            "elevator": "있음" if (cb["elevator"] or 0) > 0 else "없음",
            "approval_ym": cb["approval_ymd"].strftime("%Y/%m") if cb["approval_ymd"] else None,
            "remodel_ym": cb["remodel_ymd"].strftime("%Y/%m") if cb["remodel_ymd"] else None,
        }
        ov = overrides.get(cb["building_pk"])
        if ov:
            _apply_override(cb, ov)
        _f16_from_ymd(cb)
        cvs = value_score.compute(cb, params)
        per_area = round(cb["price"] / float(cb["total_area"]) * report_calc.M2_PER_PYEONG)
        c_la = float(cb["land_area"]) if cb["land_area"] else None
        c_gt = (float(cb["gongsi_latest"]) * c_la) if (cb["gongsi_latest"] and c_la) else None
        comps.append({"building_pk": cb["building_pk"], "addr": cb["addr"],
                      "contract_ym": cb["contract_ym"], "price": cb["price"],
                      "total_area": float(cb["total_area"]), "land_area": c_la, "gongsi_total": c_gt,
                      "score": cvs["score"],
                      "per_area": per_area, "dist_m": cb["dist_m"], "land_use": cb["land_use"],
                      "type_factor": _ADJ_FACTOR if cb["land_use"] in adj else 1.0,
                      "lng": cb["lng"], "lat": cb["lat"],
                      "fields": {**fields, **(ov or {})}, "is_outlier": False})
    _flag_comp_outliers(comps)
    return comps


async def _load_comps(building_pk: str, subject: dict, params: dict,
                      exclude: set | None = None, overrides: dict | None = None) -> list[dict]:
    """생성용: 제외 comp를 뺀 F-17 입력 리스트. exclude=None이면 이상치 기본 제외."""
    comps = await _fetch_comps(building_pk, subject, params, overrides)
    if exclude is None:
        exclude = {c["building_pk"] for c in comps if c["is_outlier"]}
    return [c for c in comps if c["building_pk"] not in exclude]


def _floor_key(fl: str) -> int:
    """층 정렬(지하=음수). '지하1층'→-1, '1층'→1."""
    import re
    m = re.search(r"(\d+)", fl or "")
    n = int(m.group(1)) if m else 0
    return -n if ("지하" in (fl or "") or (fl or "").upper().startswith("B")) else n


async def _nearby_rent_apply(building_pk: str, subject: dict, team_id: int) -> dict | None:
    """주변임대시세(반경 내 층별 평균, 이상치 제외)를 본매물 층에 적용.
    반환: {floors:[{floor,cur,mkt,diff,count}], applied_rent, applied_deposit, cur_rent, cur_deposit}
    또는 None(주변 임대사례 없음 → 토글 무의미)."""
    import statistics
    poly, radius, center = _market_spatial(subject)
    geom = await pool().fetchrow(
        "SELECT ST_X(geom) AS lng, ST_Y(geom) AS lat FROM master.buildings WHERE building_pk=$1", building_pk)
    clng = (center or {}).get("lng") if center else None
    clat = (center or {}).get("lat") if center else None
    if clng is None and geom:
        clng, clat = geom["lng"], geom["lat"]

    # 주변 평균 수익률(중앙값) = 반경 내 건물들의 (연임대추정 ÷ 적정가) — 06 비교·07 의견용.
    nearby_roi = await pool().fetchval(
        f"""SELECT percentile_cont(0.5) WITHIN GROUP (ORDER BY e.annual_rent::float / se.sale_est * 100)
            FROM master.building_rent_est e
            JOIN master.building_sale_est se USING (building_pk)
            JOIN master.buildings b USING (building_pk)
            WHERE e.annual_rent > 0 AND se.sale_est > 0 AND b.building_pk <> $4 AND {_COMP_SPATIAL}""",
        clng, clat, radius, building_pk, json.dumps(poly) if poly else None)

    # 주변 임대 comps = 팀 실제(app.floor_rents) + 마스터 추정(master.floor_rent_est, 팀 미입력 층) — market.nearby와 동일 소스.
    rows = await pool().fetch(
        f"""SELECT floor, area, rent, deposit FROM (
              SELECT fr.floor, fr.contract_area::float AS area, fr.rent::float AS rent, COALESCE(fr.deposit,0)::float AS deposit
              FROM app.floor_rents fr JOIN master.buildings b ON b.building_pk = fr.building_pk
              WHERE fr.deleted_at IS NULL AND fr.is_vacant IS NOT TRUE AND fr.building_pk <> $4
                AND fr.rent > 0 AND fr.contract_area > 0 AND {_COMP_SPATIAL}
              UNION ALL
              SELECT fo.floor, sum(fo.exclusive_area)::float, sum(fre.rent_est)::float, sum(COALESCE(fre.deposit_est,0))::float
              FROM master.buildings b JOIN master.floor_rent_est fre ON fre.building_pk = b.building_pk
                   JOIN master.floor_outline fo ON fo.building_pk = fre.building_pk AND fo.seq = fre.seq
              WHERE b.building_pk <> $4 AND fre.rent_est > 0 AND fo.exclusive_area > 0 AND {_COMP_SPATIAL}
                AND NOT EXISTS (SELECT 1 FROM app.floor_rents fr2
                                WHERE fr2.building_pk = b.building_pk AND fr2.floor = fo.floor AND fr2.deleted_at IS NULL)
              GROUP BY b.building_pk, fo.floor
            ) q WHERE area > 0 AND rent > 0""",
        clng, clat, radius, building_pk, json.dumps(poly) if poly else None)
    if not rows:
        return None
    # per_rent/per_deposit(원/㎡) → 전역 IQR 이상치 제외 → 층별 평균. market.nearby와 동일 알고리즘(화면 숫자 일치)
    recs = [{"floor": r["floor"], "per_rent": r["rent"] / r["area"],
             "per_deposit": r["deposit"] / r["area"]} for r in rows]
    prs = [x["per_rent"] for x in recs]
    lo, hi = float("-inf"), float("inf")
    if len(prs) >= 3:
        q1, q3 = statistics.quantiles(prs, n=4)[0], statistics.quantiles(prs, n=4)[2]
        lo, hi = q1 - 1.5 * (q3 - q1), q3 + 1.5 * (q3 - q1)
    by_floor: dict = {}
    for x in recs:
        if lo <= x["per_rent"] <= hi:
            by_floor.setdefault(x["floor"], []).append(x)
    mkt = {fl: {"per_rent": statistics.mean(x["per_rent"] for x in xs),
                "per_deposit": statistics.mean(x["per_deposit"] for x in xs), "count": len(xs)}
           for fl, xs in by_floor.items()}

    # 본매물 층별 = 마스터 대장(floor_outline+floor_rent_est) 기준, 팀 오버레이(app.floor_rents) 있으면 그 층 대체.
    mrows = await pool().fetch(
        """SELECT fo.floor, sum(fo.exclusive_area)::float AS area,
                  sum(fre.rent_est)::float AS rent, sum(COALESCE(fre.deposit_est,0))::float AS deposit
           FROM master.floor_outline fo JOIN master.floor_rent_est fre USING (building_pk, seq)
           WHERE fo.building_pk=$1 AND fre.rent_est>0 GROUP BY fo.floor""", building_pk)
    trows = await pool().fetch(
        """SELECT floor, sum(contract_area)::float AS area, sum(rent)::float AS rent, sum(COALESCE(deposit,0))::float AS deposit
           FROM app.floor_rents WHERE building_pk=$1 AND team_id=$2 AND deleted_at IS NULL AND is_vacant IS NOT TRUE
           GROUP BY floor""", building_pk, team_id)
    sf: dict = {r["floor"]: {"area": r["area"] or 0.0, "rent": r["rent"] or 0, "deposit": r["deposit"] or 0} for r in mrows}
    for r in trows:   # 팀 입력 층 = 대체(오버레이 우선)
        sf[r["floor"]] = {"area": r["area"] or 0.0, "rent": r["rent"] or 0, "deposit": r["deposit"] or 0}
    if not sf:
        return None
    floors, applied_rent, applied_deposit, cur_rent, cur_deposit = [], 0, 0, 0, 0
    for fl in sorted(sf, key=_floor_key):
        s = sf[fl]; m = mkt.get(fl); scur = round(s["rent"])
        cur_rent += scur; cur_deposit += round(s["deposit"])
        if m and s["area"] > 0:
            mr, md, cnt = round(m["per_rent"] * s["area"]), round(m["per_deposit"] * s["area"]), m["count"]
        else:
            mr, md, cnt = scur, round(s["deposit"]), 0   # 그 층 주변사례 없음 → 현재 폴백
        applied_rent += mr; applied_deposit += md
        floors.append({"floor": fl, "cur": scur, "mkt": mr, "diff": mr - scur, "count": cnt})
    return {"floors": floors, "applied_rent": applied_rent, "applied_deposit": applied_deposit,
            "cur_rent": cur_rent, "cur_deposit": cur_deposit,
            "nearby_roi": round(float(nearby_roi), 2) if nearby_roi else None}


def _parse_far(txt) -> float | None:
    """legal_far 텍스트('800%', '1,000% (도심 800%)') → float(%). 앞의 숫자."""
    if not txt:
        return None
    import re
    m = re.search(r"[\d,]+", str(txt))
    return float(m.group().replace(",", "")) if m else None


async def _market_zones(building_pk: str) -> list[dict]:
    """상권 존(격자 ~100m) — 셀별 지배 용도(업무/먹자/유흥/판매) 폴리곤. 지도 오버레이용."""
    rows = await pool().fetch(
        """WITH s AS (SELECT geom, ST_X(geom) lng, ST_Y(geom) lat FROM master.buildings WHERE building_pk=$1),
             cells AS (
               SELECT ST_SnapToGrid(b.geom, s.lng, s.lat, 0.0011, 0.0009) cell,
                 CASE WHEN fo.use ~ '사무소|업무시설' THEN '업무'
                      WHEN fo.use ~ '음식점' THEN '먹자'
                      WHEN fo.use ~ '유흥|단란|노래연습장|주점' THEN '유흥'
                      WHEN fo.use ~ '소매점|백화점' THEN '판매' ELSE '기타' END cat
               FROM master.buildings b JOIN master.floor_outline fo USING(building_pk), s
               WHERE ST_DWithin(b.geom::geography, s.geom::geography, 300)
                 AND fo.use !~ '주택|아파트|오피스텔|주차|부대'),
             agg AS (SELECT cell, cat, count(*) c FROM cells WHERE cat<>'기타' GROUP BY cell, cat),
             dom AS (SELECT DISTINCT ON (cell) cell, cat, c FROM agg ORDER BY cell, c DESC)
           SELECT ST_AsGeoJSON(ST_Envelope(ST_Expand(cell, 0.00055, 0.00045))) geojson, cat, c
           FROM dom WHERE c >= 3 ORDER BY c""",
        building_pk)
    return [{"geojson": json.loads(r["geojson"]), "cat": r["cat"], "count": r["c"]} for r in rows]


async def _use_type(building_pk: str, b: dict) -> dict | None:
    """F-20 활용 유형(투자 유형) 분류 — legal_far·상권 프로필 조립 후 classify()."""
    lf = await pool().fetchval(
        """SELECT max(pr.legal_far) FROM master.building_parcels bp
           JOIN master.parcels pr ON pr.pnu = bp.pnu WHERE bp.building_pk = $1""", building_pk)
    mk = await pool().fetchrow(
        """WITH s AS (SELECT geom FROM master.buildings WHERE building_pk=$1),
             f AS (SELECT fo.use FROM master.floor_outline fo JOIN master.buildings b USING(building_pk), s
                   WHERE ST_DWithin(b.geom::geography, s.geom::geography, 300)
                     AND fo.use !~ '주택|아파트|오피스텔|주차|부대|고시원')
           SELECT count(*) AS n,
             avg((use ~ '사무소|업무시설')::int)::float AS office, avg((use ~ '음식점')::int)::float AS food,
             avg((use ~ '유흥|단란|노래연습장|주점')::int)::float AS ent, avg((use ~ '소매점|백화점')::int)::float AS retail
           FROM f""", building_pk)
    market = ({"office": mk["office"] or 0, "food": mk["food"] or 0, "ent": mk["ent"] or 0, "retail": mk["retail"] or 0}
              if mk and mk["n"] else {})
    la = _fnum(b.get("land_area"))
    result = use_type.classify({
        "far": _fnum(b.get("far")), "legal_far": _parse_far(lf), "land_use": b.get("land_use"),
        "floors_above": b.get("floors_above"), "land_area_py": (la / 3.305785) if la else None,
        "age_years": b.get("age_years"), "remodel_years": b.get("remodel_years"),
        "shape": b.get("shape"), "road_frontage": b.get("road_frontage"),
        "road_score": value_score.ROAD_SCORES.get(b.get("road_frontage") or "", 0),
        "station_score": value_score.station_score(_fnum(b.get("station_dist"))),
        "use_zone": b.get("use_zone"), "market": market,
    })
    result["zones"] = await _market_zones(building_pk)   # 상권 존 폴리곤(지도용)
    return result


def synthesize(subject: dict, subject_score: float, comps: list[dict],
               params: dict, time_adjust: dict, rent_apply: dict | None = None) -> dict:
    """F-17 적정매매가 + F-18 예상수익률 + 협의금액. preview·생성 공용.
    rent_apply(토글 ON) 있으면 주변임대 적용 총임대료/보증금 사용, 없으면 현재값 폴백."""
    ap = report_calc.appraise(subject_score, subject, comps, params, time_adjust)
    if rent_apply:
        rent, deposit = rent_apply["applied_rent"], rent_apply["applied_deposit"]
    else:   # 토글 OFF or 주변사례 없음 → 현재 임대료·보증금(만실 시=공실0이라 공실제외 값)
        rent, deposit = _fnum(subject.get("total_rent")), _fnum(subject.get("total_deposit"))
    # 수익환원 블렌드(β): NOI(연임대) ÷ 구cap 을 comp식(v2)에 소폭 섞음(공용 report_calc.blend_income).
    # 임대 = 팀입력(월) 있으면 그것, 없으면 마스터 추정 연임대. 백테스트상 β=0.2가 최적.
    beta = params.get("blend.income", 0.2)
    cap = _fnum(subject.get("gu_cap"))
    ann_rent = (rent * 12) if rent else _fnum(subject.get("est_annual_rent"))
    blended = report_calc.blend_income(ap.get("fair_price"), ann_rent, cap, beta)
    if blended != ap.get("fair_price"):
        subj_py = (_fnum(subject.get("total_area")) or 0) / report_calc.M2_PER_PYEONG
        ap = {**ap, "fair_price": blended,
              "avg_per_pyeong": round(blended / subj_py) if subj_py else ap.get("avg_per_pyeong")}
    # 3층 가격(2026-07-29): 매도희망가(건물주) · 매매가(중개인, 기본=적정가) · 빌탐정 적정가(시스템=fair_price)
    ask = _fnum(subject.get("ask_price"))                            # 매도희망가 = 건물주 원하는 값(오버레이)
    broker = _fnum(subject.get("sale_price")) or ap.get("fair_price")  # 매매가 = 중개인 판단(오버레이), 없으면 적정가
    roi = report_calc.expected_roi(rent, broker or ap["fair_price"])   # 수익률은 실제 매수기준가(매매가)로
    gap = round(ask - broker) if (ask and broker) else None          # 협의금액 = 매도희망가 − 매매가
    if ap.get("breakdown"):   # 수익환원 블렌드 정보 보강(리빌용)
        ap["breakdown"] = {**ap["breakdown"], "beta": round(beta, 2),
                           "income_val": round(ann_rent / cap) if (ann_rent and cap) else None,
                           "final": ap.get("fair_price")}
    # 공시지가 맥락(05 페이지): 주변 사례 공시지가 중앙값(원/㎡) + 공시배율(실거래÷공시총액)
    import statistics as _st
    _pm2 = [c["gongsi_total"] / c["land_area"] for c in comps if c.get("gongsi_total") and c.get("land_area")]
    _mult = [c["price"] / c["gongsi_total"] for c in comps if c.get("gongsi_total") and c.get("price")]
    gongsi_ctx = {"nbhd_per_m2": round(_st.median(_pm2)) if _pm2 else None,
                  "mult": round(_st.median(_mult), 2) if _mult else None, "n": len(_pm2)}
    # 임대 요약(06 페이지): 층수·현재/주변 총임대료·보증금 집계
    rent_summary = ({"floor_count": len(rent_apply["floors"]),
                     "cur_rent": round(rent_apply["cur_rent"]), "mkt_rent": round(rent_apply["applied_rent"]),
                     "cur_deposit": round(rent_apply["cur_deposit"]), "mkt_deposit": round(rent_apply["applied_deposit"]),
                     "nearby_roi": rent_apply.get("nearby_roi")}
                    if rent_apply else None)
    return {**ap, "expected_roi": roi, "gap": gap, "gongsi_ctx": gongsi_ctx, "rent_summary": rent_summary,
            "ask_price": round(ask) if ask else None,               # 매도희망가
            "broker_price": round(broker) if broker else None,     # 매매가(중개인)
            "applied_rent": round(rent) if rent else None, "expected_deposit": round(deposit) if deposit else None,
            "rent_floors": rent_apply["floors"] if rent_apply else None, "market_applied": bool(rent_apply)}


def _outlier_bounds(vals: list[float]) -> tuple[float, float] | None:
    """평단가 이상치 [하한,상한]. 표본 충분(≥10)=IQR 1.5, 소표본=MAD 수정z(3.5, 소표본서도 견고).
    3건 미만이거나 편차 0이면 판정 불가(None). 배치·라이브 공용 규칙."""
    import statistics
    if len(vals) < 3:
        return None
    med = statistics.median(vals)
    if len(vals) >= 10:
        q1, q3 = statistics.quantiles(vals, n=4)[0], statistics.quantiles(vals, n=4)[2]
        iqr = q3 - q1
        return (q1 - 1.5 * iqr, q3 + 1.5 * iqr)
    mad = statistics.median([abs(x - med) for x in vals])
    if mad == 0:
        return None
    d = 3.5 * mad / 0.6745   # Iglewicz-Hoaglin 수정 z-score 임계 3.5
    return (med - d, med + d)


def _flag_comp_outliers(comps: list[dict]) -> None:
    """평단가 이상치 플래그(F-17 기본 제외)."""
    b = _outlier_bounds([c["per_area"] for c in comps])
    if b is None:
        return
    lo, hi = b
    for c in comps:
        if not (lo <= c["per_area"] <= hi):
            c["is_outlier"] = True


TEMPLATE_ANALYSIS = os.path.join(os.path.dirname(__file__), "../../../specs/03-features/R_example.pptx")

# ── 템플릿 바인딩 헬퍼 ──────────────────────────────────────────────
def _run0(shape, text: str) -> None:
    """첫 문단 첫 run 텍스트만 교체(큰 숫자 표시용, 서식·단위 run 유지)."""
    p = shape.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
    else:
        p.text = text


def _settext(shape, text: str) -> None:
    """도형 전체 텍스트를 한 줄로 교체(서술 중화용). 첫 run 서식 유지, 나머지 run 비움."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.text = text
    for p in tf.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


def _cell(cell, text: str) -> None:
    p = cell.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


def _eok(won, dec=0):
    return f"{won/1e8:.{dec}f}" if won else "—"


def _man(won):
    return f"{round(won/1e4):,}" if won else "—"


def _eokman(won):
    if not won:
        return "—"
    e = int(won // 1e8); m = round((won % 1e8) / 1e4)
    return (f"{e}억 " if e else "") + f"{m:,}만원"


def _py(m2):
    return f"{m2/3.305785:.2f}" if m2 else "—"


def _grade_word(sc):
    return "매우 우수" if sc >= 90 else "우수" if sc >= 80 else "양호" if sc >= 70 else "보통" if sc >= 60 else "미흡"


def _bind_analysis_template(path: str, report_id: int, b: dict, vs: dict, syn: dict) -> int:
    """R_example.pptx 서식에 실 산출값 바인딩. 파생 불가한 매물특정 서술은 중화(허위 방지)."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData

    prs = Presentation(TEMPLATE_ANALYSIS)
    fair = syn.get("fair_price"); ask = syn.get("ask_price"); gap = syn.get("gap")
    rent = syn.get("applied_rent") or _fnum(b.get("total_rent"))
    dep = syn.get("expected_deposit") or _fnum(b.get("total_deposit"))
    cur_rent = _fnum(b.get("total_rent"))
    roi = syn.get("expected_roi")
    roi_ask = report_calc.expected_roi(rent, ask) if ask else None
    area_m2 = _fnum(b.get("total_area"))
    floors = syn.get("rent_floors") or []
    used = syn.get("comps_used") or []
    items = vs["items"]
    today = dt.date.today().strftime("%Y.%m.%d")
    rno = f"BT-{dt.date.today().year}-{report_id:06d}"
    addr = b.get("addr") or ""

    # 1) 단어 통일 + 매물특정 예시값 치환(모든 run, 표·차트 캐시 포함). 긴 토큰 우선.
    SUBS = [
        ("주변월세시세", "주변임대시세"), ("총월세", "총임대료"), ("월세", "임대료"),
        ("역삼동 735-29", addr), ("BT-2026-000104", rno), ("2026.06.28", today),
        ("1억 9,716만원", _eokman(rent * 12)),
        ("A등급(양호)", f"{vs['grade']}등급({_grade_word(vs['score'])})"),
    ]
    def _subs_runs(paras):
        for para in paras:
            for r in para.runs:
                for ex, real in SUBS:
                    if ex in r.text:
                        r.text = r.text.replace(ex, real)

    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                _subs_runs(sh.text_frame.paragraphs)
            elif sh.has_table:
                for row in sh.table.rows:
                    for cl in row.cells:
                        _subs_runs(cl.text_frame.paragraphs)

    # 2) 슬라이드별 스칼라·표·차트 바인딩. 도형은 이름('Text N')으로 지정(슬라이드 내 유일).
    def sid(slide, n):
        return next((sh for sh in slide.shapes if sh.name == f"Text {n}"), None)

    def tbl(slide):
        return next((sh.table for sh in slide.shapes if sh.has_table), None)

    def chart(slide):
        return next((sh.chart for sh in slide.shapes if sh.has_chart), None)

    S = prs.slides
    # ── 표지(0) ──
    _run0(sid(S[0], 2), addr)
    _settext(sid(S[0], 3), f"{addr} 분석보고서")
    _run0(sid(S[0], 8), f"{vs['score']:.1f}")
    _run0(sid(S[0], 7), f"{vs['grade']}등급")
    _run0(sid(S[0], 11), _eok(ask))
    _run0(sid(S[0], 14), _eok(fair))
    _run0(sid(S[0], 19), _eok(gap) if gap else "—")
    _settext(sid(S[0], 21), f"주변임대시세 적용 총임대료 · 연임대료 추정치 {_eokman(rent*12)}")
    _run0(sid(S[0], 22), _man(rent))
    _run0(sid(S[0], 25), f"{roi}" if roi is not None else "—")

    # ── 기본정보(1) ──
    def _d(x):
        return "—" if x in (None, "") else x
    land_py = (_fnum(b.get("land_area")) or 0) / 3.305785
    total_py = (area_m2 or 0) / 3.305785
    t = tbl(S[1])
    binfo = [_eok(ask) + "억 원", f"{_py(_fnum(b.get('land_area')))}평", f"{_py(area_m2)}평",
             (_eok(ask/land_py, 2) + "억 원 (호가 기준)") if (ask and land_py) else "—",
             (_man(ask/total_py) + "만 원 (호가 기준)") if (ask and total_py) else "—",
             _d(b.get("use_zone")), _d(b.get("main_use")),
             f"지하 {_d(b.get('floors_below'))}층 / 지상 {_d(b.get('floors_above'))}층",
             (str(b.get("approval_ymd"))[:4] + "년") if b.get("approval_ymd") else "—",
             f"{_d(b.get('bcr'))}%", f"{_d(b.get('far'))}%",
             f"{_d(b.get('parking_count'))}대", f"{_d(b.get('elevator_count'))}대",
             "—", f"공실 {b.get('vacant_count',0)}건",
             _eok(_fnum(b.get("total_deposit")), 1) + "억 원", _man(cur_rent) + "만 원", "—",
             (f"{report_calc.expected_roi(cur_rent, ask)}% (현재 임대료 · 호가 기준)" if ask else "—")]
    for i, v in enumerate(binfo):
        if i < len(t.rows):
            _cell(t.cell(i, 1), v)
    # 핵심스펙 카드: 파생 가능만 채우고 매물특정(초역세권·코너)은 중화
    _settext(sid(S[1], 8), "용도지역"); _settext(sid(S[1], 9), _d(b.get("use_zone")))
    _settext(sid(S[1], 11), "건물규모"); _settext(sid(S[1], 12), f"지하 {_d(b.get('floors_below'))} · 지상 {_d(b.get('floors_above'))}층")
    _settext(sid(S[1], 15), f"승강기 {_d(b.get('elevator_count'))}대 보유")
    _settext(sid(S[1], 18), f"주차 {_d(b.get('parking_count'))}대 가능")
    _settext(sid(S[1], 21), f"현재 {report_calc.expected_roi(cur_rent, ask) if ask else '—'}%  →  적정가 기준 {roi if roi is not None else '—'}%")
    _settext(sid(S[1], 23), (f"현재 총임대료 {_man(cur_rent)}만원 대비 주변임대시세 적용 시 "
                             f"{_man(rent)}만원({(rent-cur_rent)/1e4:+,.0f}만원)까지 임대료 개선 여지가 있습니다."
                             if syn.get("market_applied") else "주변임대시세 미포함(현재 임대료 기준)."))

    # ── 가치점수(3) ──
    _run0(sid(S[3], 8), f"{vs['score']:.1f}")
    _run0(sid(S[3], 10), vs["grade"])
    t = tbl(S[3])
    ORDER = ["road_access", "station_dist", "use_zone", "approval_date", "elevator", "remodel", "shape", "slope", "float_pop"]
    for i, key in enumerate(ORDER, start=1):
        sc = items.get(key, 0)
        _cell(t.cell(i, 1), _grade_word(sc))
        _cell(t.cell(i, 2), f"{t.cell(i,0).text.strip()} 항목 평가 결과 {_grade_word(sc)} 수준입니다.")
    ch = chart(S[3])
    if ch:
        cd = CategoryChartData()
        cd.categories = ["도로접면", "역과의거리", "용도지역", "지형형상", "사용승인일", "엘리베이터", "대수선·리모델링", "경사도", "유동인구"]
        cd.add_series("항목별 점수(배점 대비 %)",
                      [items.get(k, 0) for k in ["road_access", "station_dist", "use_zone", "shape", "approval_date", "elevator", "remodel", "slope", "float_pop"]])
        ch.replace_data(cd)
    _settext(sid(S[3], 11), f"등급 기준  S 90↑ · A 75~89 · B 60~74 · C 60↓   →   본 매물 {vs['grade']}등급({_grade_word(vs['score'])})")

    # ── 매매사례(4) ──
    t = tbl(S[4])
    for i in range(1, len(t.rows)):
        c = used[i - 1] if i - 1 < len(used) else None
        vals = ([str(i), c.get("addr") or c.get("building_pk") or "—", "—", str(c.get("contract_ym") or "—"),
                 f"{c['price']/1e8:.1f}", f"{c.get('area_py','—')}", f"{round(c['per_now']/1e4):,}",
                 f"{vs['score']-c.get('score',0):+.1f}", f"{round(c.get('time_adj',0)*100):+d}%"]
                if c else [str(i)] + ["—"] * 8)
        for j, v in enumerate(vals):
            _cell(t.cell(i, j), v)
    ch = chart(S[4])
    if ch and used:
        cd = CategoryChartData()
        cd.categories = [f"사례{i+1}" for i in range(len(used))] + ["본 매물"]
        cd.add_series("유사사례 평단가", [round(c["per_now"] / 1e4) for c in used] + [None])
        cd.add_series("본 매물 적용", [None] * len(used) + [round((syn.get("avg_per_pyeong") or 0) / 1e4)])
        ch.replace_data(cd)
    _settext(sid(S[4], 9), f"{_man(syn.get('avg_per_pyeong'))} 만원/평")
    _settext(sid(S[4], 13), f"{_py(area_m2)} 평")
    _settext(sid(S[4], 17), f"{_eok(fair)}억 원")
    if gap and ask:
        _settext(sid(S[4], 20), f"현재 호가 {_eok(ask)}억 대비 약 {_eok(gap)}억 협의 필요 (호가 대비 약 {gap/ask*100:.1f}% 하향 협의 여지)")
    else:
        _settext(sid(S[4], 20), "적정매매가 수준 · 협의 여지 제한적")

    # ── 주변임대시세(5) ──
    _run0(sid(S[5], 8), _man(cur_rent))
    _settext(sid(S[5], 11), f"{_man(rent)} 만원    연임대료 추정치 {_eokman(rent*12)}")
    _run0(sid(S[5], 14), f"{(rent-cur_rent)/1e4:+,.0f}")
    t = tbl(S[5])
    n = min(len(floors), len(t.rows) - 2)   # 마지막 행=합계
    for i in range(len(t.rows) - 2):        # 데이터 행: 없는 층은 비움
        if i < n:
            f = floors[i]
            for j, v in enumerate([f["floor"], f"{_man(f['cur'])}만원", f"{_man(f['mkt'])}만원",
                                   f"{f['diff']/1e4:+,.0f}만원", f"{f['count']}건"]):
                _cell(t.cell(i + 1, j), v)
        else:
            for j in range(5):
                _cell(t.cell(i + 1, j), "")
    last = len(t.rows) - 1
    for j, v in enumerate(["합계", f"{_man(cur_rent)}만원", f"{_man(rent)}만원",
                           f"{(rent-cur_rent)/1e4:+,.0f}만원", f"{sum(f['count'] for f in floors)}건"]):
        _cell(t.cell(last, j), v)
    ch = chart(S[5])
    if ch and floors:
        cd = CategoryChartData()
        cd.categories = [f["floor"] for f in floors]
        cd.add_series("현재 임대료", [round(f["cur"] / 1e4) for f in floors])
        cd.add_series("주변임대시세", [round(f["mkt"] / 1e4) for f in floors])
        ch.replace_data(cd)

    # ── 예상수익률(6) ──
    _run0(sid(S[6], 8), _eok(dep, 1))
    _run0(sid(S[6], 11), _man(rent))
    _run0(sid(S[6], 14), _eokman(rent * 12).replace("만원", ""))
    _run0(sid(S[6], 17), f"{roi}" if roi is not None else "—")
    t = tbl(S[6])
    _cell(t.cell(1, 1), f"{_eok(ask)}억 원"); _cell(t.cell(1, 2), f"{_eok(fair)}억 원")
    _cell(t.cell(2, 1), f"{roi_ask}%" if roi_ask is not None else "—")
    _cell(t.cell(2, 2), f"{roi}%" if roi is not None else "—")
    _settext(sid(S[6], 23), f"{_man(rent)}만원")
    _settext(sid(S[6], 31), _eokman(rent * 12))
    _settext(sid(S[6], 35), f"{_eok(fair)}억 원")
    _settext(sid(S[6], 39), f"{roi}%" if roi is not None else "—")
    _settext(sid(S[6], 42), (f"주변임대시세 적용 총임대료 {_man(rent)}만원 기준 · 예상보증금 {_eok(dep,1)}억 원 반영 · "
                             f"적정매매가 {_eok(fair)}억 원 기준 단순 연임대수익 산정"))

    # ── 최종요약(7) ──
    _run0(sid(S[7], 9), _eok(ask))
    _run0(sid(S[7], 14), _eok(fair))
    _run0(sid(S[7], 19), _eok(gap) if gap else "—")
    _settext(sid(S[7], 22), (f"유사 매매사례 분석 결과, 본 매물의 적정매매가는 약 {_eok(fair)}억 원 수준입니다. "
                             + (f"현재 매도희망가 {_eok(ask)}억 원은 적정가 대비 약 {_eok(gap)}억 원 높은 수준입니다." if gap and gap > 0 else "")))
    _settext(sid(S[7], 25), "적정매매가 기준 접근 필요 · 주변임대시세 적용 시 수익성 개선 가능 · 가격 협의 여부가 투자 판단의 핵심")
    _run0(sid(S[7], 29), _eok(dep, 1) + " 억원")
    _run0(sid(S[7], 32), _man(rent) + " 만원")
    _run0(sid(S[7], 35), _eokman(rent * 12).replace("만원", " 만원"))
    _run0(sid(S[7], 38), (f"{roi} %" if roi is not None else "—"))

    prs.save(path)
    return len(prs.slides.__iter__.__self__._sldIdLst)  # noqa: SLF001


def _make_pptx(path: str, kind: str, b: dict, vs: dict | None, syn: dict | None = None,
               report_id: int = 0) -> int:
    """python-pptx로 보고서 생성. 반환=슬라이드 수.
    analysis = R_example.pptx 서식 바인딩 / briefing = 텍스트 슬라이드(전용 서식 없음)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    if kind == "analysis" and vs is not None and os.path.exists(TEMPLATE_ANALYSIS):
        return _bind_analysis_template(path, report_id, b, vs, syn or {})

    prs = Presentation()
    blank = prs.slide_layouts[6]

    def slide(title: str, lines: list[str]):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(1))
        tb.text_frame.text = title
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        body = s.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(9), Inches(5))
        tf = body.text_frame
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ln
            p.font.size = Pt(14)

    addr = b.get("addr", "")
    if kind == "briefing":  # 7슬라이드(R §4)
        slide("빌탐정 브리핑 자료", [addr])
        slide("매물 기본정보", [
            f"대지면적 {b.get('land_area','—')}㎡ · 연면적 {b.get('total_area','—')}㎡",
            f"층수 지상{b.get('floors_above','—')}/지하{b.get('floors_below','—')} · 용적률 {b.get('far','—')}%",
        ])
        slide("위치 · 지도", ["(Static Map — 프로덕션 연동)"])
        slide("로드뷰 / 사진", ["(Panorama — 프로덕션 연동)"])
        slide("임대 내역", [
            f"총보증금 {b['total_deposit']:,}원 · 총임대료 {b['total_rent']:,}원 · 공실 {b['vacant_count']}건",
        ])
        slide("추가 사진", ["—"])
        slide("마무리", ["빌탐정 BILLTAMJUNG"])
    else:  # analysis 8슬라이드(R §5)
        assert vs is not None
        slide("매물분석보고서", [addr, f"가치점수 {vs['score']} · {vs['grade']}등급"])
        slide("매물 기본정보", [
            f"대지 {b.get('land_area','—')}㎡ · 연면적 {b.get('total_area','—')}㎡ · 용적률 {b.get('far','—')}%",
        ])
        slide("분석 흐름", ["STEP1 가치점수 → STEP2 매매사례 → STEP3 주변임대 → STEP4 수익률"])
        slide("STEP1 가치점수", [
            f"총점 {vs['score']} / 100 · {vs['grade']}등급",
            *[f"{k}: {v}" for k, v in vs["items"].items()],
        ])
        syn = syn or {}
        used = syn.get("comps_used") or []
        fair = syn.get("fair_price")
        slide("STEP2 매매사례 시세분석", [
            f"유효 사례 {len(used)}건 · 가중평균 평단가 {syn.get('avg_per_pyeong') or '—'}원/평",
            *[f"{u.get('addr') or u['building_pk']} · {u['contract_ym']} · "
              f"{u['price']:,}원 · 가치 {u['score']} · 가중 {u['weight']}" for u in used[:12]],
        ])
        if syn.get("market_applied"):
            step3 = [f"적용 총임대료 {syn['applied_rent']:,}원 (주변 임대시세 반영)"]
            step3 += [f"{f['floor']}: 현재 {f['cur']:,} → 주변 {f['mkt']:,} (차이 {f['diff']:+,}, 사례 {f['count']}건)"
                      for f in (syn.get("rent_floors") or [])]
        else:
            step3 = [f"적용 총임대료 {b['total_rent']:,}원 (현재 임대 기준 · 주변시세 제외)"]
        slide("STEP3 주변임대시세", step3)
        slide("STEP4 적정매매가·예상수익률", [
            f"적정매매가 {fair:,}원" if fair else "적정매매가 — (유효 매매사례 없음)",
            f"예상수익률 {syn.get('expected_roi')}%" if syn.get("expected_roi") is not None else "예상수익률 —",
            f"협의 필요금액 {syn.get('gap'):,}원" if syn.get("gap") is not None else "협의 필요금액 —",
        ])
        slide("최종 요약", [
            f"가치점수 {vs['score']}({vs['grade']}) · 적정매매가 {fair:,}원" if fair
            else f"가치점수 {vs['score']}({vs['grade']})",
            "빌탐정 BILLTAMJUNG",
        ])

    prs.save(path)
    return len(prs.slides.__iter__.__self__._sldIdLst)  # noqa: SLF001


async def run_generate(report_id: int, team_id: int) -> dict:
    """잡 본체. 성공=크레딧 차감+완료 / 실패=failed+미차감."""
    os.makedirs(REPORT_DIR, exist_ok=True)
    try:
        async with tx() as conn:
            rep = await conn.fetchrow("SELECT * FROM app.reports WHERE id=$1", report_id)
            if not rep:
                return {"ok": False, "reason": "not found"}
            await conn.execute("UPDATE app.reports SET status='generating' WHERE id=$1", report_id)

        fs_version, params, time_adjust = await _load_formula_params()
        b = await _assemble(rep["building_pk"], team_id)
        vs = value_score.compute(b, params) if rep["kind"] == "analysis" else None

        syn = None
        if rep["kind"] == "analysis":   # F-17 적정매매가 · F-18 예상수익률
            opt = rep["options_json"]
            opt = json.loads(opt) if isinstance(opt, str) else (opt or {})
            exclude = set(opt.get("exclude") or [])
            overrides = opt.get("overrides") or {}
            comps = await _load_comps(rep["building_pk"], b, params, exclude, overrides)
            rent_apply = None
            if opt.get("include_market", True):   # 토글 ON → 주변임대 적용(STEP3·F-18)
                rent_apply = await _nearby_rent_apply(rep["building_pk"], b, team_id)
            syn = synthesize(b, vs["score"], comps, params, time_adjust, rent_apply)

        path = os.path.join(REPORT_DIR, f"report_{report_id}.pptx")
        _make_pptx(path, rep["kind"], b, vs, syn, report_id)

        # 웹 보고서(/reports/:id) 렌더용 synthesis 스냅샷 — 생성 시점 값 고정(analysis만).
        snapshot = None
        if rep["kind"] == "analysis" and vs and syn:
            ut = await _use_type(rep["building_pk"], b)   # F-20 투자 유형
            snapshot = {
                "subject": {"addr": b.get("addr"), "score": vs["score"], "grade": vs["grade"],
                            "items": vs["items"], "total_area": _fnum(b.get("total_area")),
                            "land_area": _fnum(b.get("land_area")), "sale_price": _fnum(b.get("sale_price")),
                            "total_rent": _fnum(b.get("total_rent"))},
                "preview": {"score": vs["score"], "grade": vs["grade"], "fair_price": syn["fair_price"],
                            "avg_per_pyeong": syn["avg_per_pyeong"], "expected_roi": syn["expected_roi"],
                            "gap": syn["gap"], "ask_price": syn["ask_price"], "broker_price": syn.get("broker_price"),
                            "applied_rent": syn.get("applied_rent"), "expected_deposit": syn.get("expected_deposit"),
                            "market_applied": syn.get("market_applied", False), "breakdown": syn.get("breakdown"),
                            "gongsi_ctx": syn.get("gongsi_ctx"), "use_type": ut,
                            "rent_summary": syn.get("rent_summary"),
                            "rent_floors": syn.get("rent_floors"), "comps_used": syn.get("comps_used")},
            }

        cost = settings.cost_analysis if rep["kind"] == "analysis" else settings.cost_briefing
        async with tx() as conn:  # 성공 트랜잭션: 차감+완료+워터마크 원자
            await conn.execute("SELECT app.deduct_credit($1,$2,$3)", rep["account_id"], cost, report_id)
            await conn.execute(
                """UPDATE app.reports SET status='done', completed_at=now(),
                     credits_spent=$2, file_path=$3, formula_set_version=$4, result_json=$7,
                     master_version=(SELECT version FROM master.master_version),
                     source_watermark=COALESCE(app.building_watermark($5,$6), now())
                   WHERE id=$1""",
                report_id, cost, path, fs_version, rep["building_pk"], team_id,
                json.dumps(snapshot) if snapshot else None,
            )
        return {"ok": True, "file": path, "credits": cost}
    except Exception as e:  # 실패: 미차감
        async with tx() as conn:
            await conn.execute(
                "UPDATE app.reports SET status='failed', failed_reason=$2 WHERE id=$1",
                report_id, str(e)[:500],
            )
        return {"ok": False, "reason": str(e)}


class GenerateIn(BaseModel):
    report_id: int
    team_id: int


@router.post("/generate-report")
async def generate_report(body: GenerateIn):
    """Cloud Tasks 진입점(프로덕션)."""
    return await run_generate(body.report_id, body.team_id)
