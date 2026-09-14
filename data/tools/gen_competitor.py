#!/usr/bin/env python3
"""경쟁사 분석 PPT — 3개 서비스별 독립 덱 생성.
  · 부기사.pptx   (유료 회원제 → 공개화면 기준 표면 분석 · 부기맨 상업 라인 포함)
  · 디스코.pptx   (Playwright로 직접 조작, 역삼동 601-5 실사용)
  · 건물닷컴.pptx (가격평가 미리보기 리포트 직접 열람)
객관적 프레이밍(빌탐정 우위 강변 배제).

실행: python gen_competitor.py            → 3개 덱 모두
      python gen_competitor.py bugisa     → 부기사만 (디스코/건물닷컴 스크린샷 없어도 OK)
      python gen_competitor.py disco gunmul
"""
import sys
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

SHOTS="/private/tmp/claude-501/-Users-iseung-ug-Desktop---------/ddcd61b1-bac0-48e0-b3bb-bca11bf8983c/scratchpad/shots/"
OUT="specs/05-business/"
AC=RGBColor(0x1E,0x5A,0xF0); INK=RGBColor(0x0F,0x1A,0x2E); SUB=RGBColor(0x3A,0x46,0x57)  # 디자인토큰 2026-07-21
LINE=RGBColor(0xE6,0xEA,0xEF); CARD=RGBColor(0xF5,0xF8,0xFC); WHITE=RGBColor(0xFF,0xFF,0xFF)
GREEN=RGBColor(0x2A,0x7D,0x46); AMBER=RGBColor(0xC8,0x78,0x1A); RED=RGBColor(0xC0,0x39,0x2B)
NAVY=RGBColor(0x0F,0x1A,0x2E)
W=13.333

def deck():
    p=Presentation(); p.slide_width=I(13.333); p.slide_height=I(7.5)
    return p
def slide(p): return p.slides.add_slide(p.slide_layouts[6])
def rect(s,l,t,w,h,fill,ln=None):
    sh=s.shapes.add_shape(1,I(l),I(t),I(w),I(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.line.color.rgb=ln if ln else fill; sh.line.width=Pt(0.75 if ln else 0); sh.shadow.inherit=False
    if not ln: sh.line.fill.background()
    return sh
def txt(s,l,t,w,h,runs,size=13,bold=False,color=INK,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sp=None):
    tb=s.shapes.add_textbox(I(l),I(t),I(w),I(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=anchor
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,0)
    if isinstance(runs,str): runs=[[(runs,size,bold,color)]]
    if runs and isinstance(runs[0],tuple): runs=[runs]
    for i,para in enumerate(runs):
        pa=tf.paragraphs[0] if i==0 else tf.add_paragraph(); pa.alignment=align
        if sp: pa.space_after=Pt(sp)
        for (t_,sz,bd,cl) in para:
            r=pa.add_run(); r.text=t_; r.font.size=Pt(sz); r.font.bold=bd; r.font.color.rgb=cl; r.font.name='Malgun Gothic'
    return tb
def pic(s,path,l,t,boxw,boxh):
    im=Image.open(path); iw,ih=im.size; ar=iw/ih
    w,h=boxw,boxw/ar
    if h>boxh: h=boxh; w=boxh*ar
    x=l+(boxw-w)/2
    pc=s.shapes.add_picture(path,I(x),I(t),I(w),I(h))
    pc.line.color.rgb=LINE; pc.line.width=Pt(0.75)
def band(s,name,accent,label):
    rect(s,0,0,W,0.6,NAVY); rect(s,0,0,0.16,0.6,accent)
    txt(s,0.42,0.15,8,0.3,[(f"경쟁사 분석  ·  {name}",12,True,WHITE)])
    txt(s,W-4.5,0.17,4,0.3,label,10,False,RGBColor(0xB8,0xC4,0xDE),PP_ALIGN.RIGHT)
def foot(s,n): txt(s,W-2.0,7.12,1.6,0.25,f"{n}",9,False,RGBColor(0xAA,0xB2,0xBC),PP_ALIGN.RIGHT)
def usedchip(s,l,t,used=True):
    w=1.55 if used else 2.1; c=GREEN if used else AMBER; lab="직접 사용" if used else "유료 · 표면 분석"
    rect(s,l,t,w,0.34,c); txt(s,l,t+0.055,w,0.3,lab,10,True,WHITE,PP_ALIGN.CENTER)
def bullets(s,l,t,w,items,size=12,gap=5):
    runs=[[("• ",size,True,AC),(a,size,b,INK)]+([(" — "+c,size-1,False,SUB)] if c else []) for (a,b,c) in items]
    txt(s,l,t,w,5,runs,size,sp=gap)
def cover(p,name,eng,tag,accent,used):
    s=slide(p); rect(s,0,0,W,7.5,NAVY); rect(s,0,0,0.35,7.5,accent)
    txt(s,0.95,1.7,11,0.35,"경쟁사 분석",16,True,RGBColor(0x8F,0xB0,0xFF))
    txt(s,0.9,2.35,11,0.8,name,44,True,WHITE)
    txt(s,0.95,3.5,11,0.4,eng,16,False,RGBColor(0xB8,0xC4,0xDE))
    txt(s,0.95,4.35,11.5,0.4,tag,15,False,RGBColor(0xC8,0xD4,0xEE))
    usedchip(s,0.95,5.25,used)
    txt(s,0.95,6.6,11.5,0.3,"시장 조사 자료 · 2026.07 · 테스트 물건: 강남구 역삼동 601-5",11,False,RGBColor(0x9A,0xA8,0xC8))
def closing(s,name,accent,gaps,verdict):
    txt(s,0.5,0.85,10,0.3,"종합 평가",12,True,accent); txt(s,0.5,1.18,11,0.5,f"{name}의 한계",22,True,INK)
    for i,(k,v) in enumerate(gaps):
        t=2.1+i*1.05; rect(s,0.5,t,12.3,0.92,CARD,LINE); rect(s,0.5,t,0.1,0.92,accent)
        txt(s,0.85,t+0.13,3.3,0.3,k,13.5,True,INK); txt(s,4.2,t+0.16,8.4,0.6,v,12,False,SUB)
    rect(s,0.5,2.1+len(gaps)*1.05+0.15,12.3,0.95,RGBColor(0xEF,0xF3,0xFB),accent)
    txt(s,0.85,2.1+len(gaps)*1.05+0.3,11.8,0.7,[[("총평  ",12.5,True,accent),(verdict,12.5,False,INK)]])

# ═══════════════════ 1. 부기사 (5장) ═══════════════════
def build_bugisa():
    p=deck(); AC1=RGBColor(0x2E,0x6F,0xED)
    cover(p,"부기사","Bugisa · 부기맨  —  중개 운영 올인원(CRM)",
          "매물·고객·직원·정산 + 마케팅·AI 올인원. 상업 빌딩은 '부기맨' 라인.",AC1,used=False)

    # ── 01 · 운영 올인원 (CRM 코어) ──
    s=slide(p); band(s,"부기사",AC1,"01 · 운영 기능(올인원)")
    txt(s,0.5,0.8,8,0.3,"중개 운영 올인원 — 방대한 기능 폭",12,True,AC1)
    txt(s,0.5,1.13,8.5,0.4,"매물·고객·직원·정산·마케팅을 한 곳에서",19,True,INK)
    usedchip(s,10.8,0.9,False)
    cats=[
     ("매물관리","주소→대장·실거래가·승강기 자동입력 · 지도기반 현장관리 · 건물 노후도 · 필지 합산(복합 브리핑)"),
     ("고객 CRM","조건 저장→매칭 매물 자동 알림 · 고객별 TM 관리"),
     ("직원·정산","활동추적(TM·계약·사진) · 권한/매물보안 · 수수료 배분·급여 산출"),
     ("마케팅·AI","1클릭 네이버부동산·SNS 발송 · 반복 포스팅 · 제로콜(AI 전화응대·통화요약)"),
    ]
    for i,(k,v) in enumerate(cats):
        t=1.8+i*1.08
        rect(s,0.5,t,6.3,0.96,CARD,LINE); rect(s,0.5,t,0.1,0.96,AC1)
        txt(s,0.78,t+0.12,5.85,0.3,k,13,True,INK)
        txt(s,0.78,t+0.44,5.85,0.5,v,10.5,False,SUB)
    pic(s,SHOTS+"bg_maemul.png",7.0,1.7,5.9,1.9)
    txt(s,7.0,3.55,5.9,0.25,"▲ 매물관리 그리드+모바일 — 주소 입력→대장·실거래가 자동",9,False,SUB)
    pic(s,SHOTS+"bg_crm.png",7.0,4.0,5.9,1.9)
    txt(s,7.0,5.82,5.9,0.25,"▲ 고객 TM관리·매물매칭 — 조건 맞는 매물 등록 시 자동 알림",9,False,SUB)
    foot(s,2)

    # ── 02 · 부기맨(상업 빌딩 특화) ──
    s=slide(p); band(s,"부기사",AC1,"02 · 부기맨(상업 빌딩 특화)")
    txt(s,0.5,0.8,9,0.3,"부기맨 — 상업용·빌딩 특화 라인",12,True,AC1)
    txt(s,0.5,1.13,9,0.4,"우리와 직접 겹치는 세그먼트",19,True,INK)
    bullets(s,0.5,1.85,6.3,[
     ("포지셔닝",True,"'건물 매매·임대 전문' 최적화 매물관리 — 별도 브랜드(bugiman.co.kr)"),
     ("상업 특화",True,"상업용 템플릿(건물개요 브리핑 프린트)·권리분석·지도기반·ChatGPT 표방"),
     ("데이터",True,"건축물/토지대장 자동입력 · PC·모바일 최적화"),
     ("가격",True,"초기 400만 + 연 100만 갱신 — 일반 대비 고가"),
     ("규모",True,"1,000+ 가입사(국내 10만 공인중개사 대상)"),
    ],size=12,gap=8)
    rect(s,0.5,5.35,6.3,1.35,RGBColor(0xEF,0xF3,0xFB),AC1)
    txt(s,0.72,5.5,5.9,1.1,[[("의미  ",11.5,True,AC1),("상업 빌딩 세그먼트를 이미 겨냥한 유일한 경쟁 라인. 매물관리·브리핑까지는 커버 — 심화 가치평가는 다음 장.",11,False,INK)]],11)
    pic(s,SHOTS+"bm_hero.png",7.0,1.55,5.9,2.45)
    pic(s,SHOTS+"bm_app.png",7.0,4.25,5.9,1.35)
    txt(s,7.0,5.62,5.9,0.3,"▲ 부기맨 상업용 매물관리 — 건물/토지대장 자동입력·화면 최적화",9,False,SUB)
    foot(s,3)

    # ── 03 · 데이터·분석과 산출물 (정정) ──
    s=slide(p); band(s,"부기사",AC1,"03 · 데이터·분석과 산출물")
    txt(s,0.5,0.8,10,0.3,"데이터·분석 — 어디까지 하나",12,True,AC1)
    txt(s,0.5,1.13,10,0.4,"'제공하는 것'과 '하지 않는 것'을 정확히",18.5,True,INK)
    rect(s,0.5,1.85,6.15,2.55,RGBColor(0xEF,0xF6,0xEF),GREEN)
    txt(s,0.72,2.0,5.7,0.3,"제공하는 것",12.5,True,GREEN)
    bullets(s,0.72,2.42,5.75,[
     ("자동 데이터",True,"건축물·토지대장·실거래가·승강기·고시정보(신규/변경/폐지)"),
     ("등기부 권리분석",True,""),
     ("부동산 계산기",True,"평단가·수익률·대출"),
     ("브리핑 산출물",True,"건물개요 IM/팜플렛 프린트(상업 템플릿)"),
    ],size=11,gap=6)
    rect(s,0.5,4.55,6.15,2.15,RGBColor(0xFB,0xEE,0xEC),RED)
    txt(s,0.72,4.7,5.7,0.3,"하지 않는 것",12.5,True,RED)
    bullets(s,0.72,5.12,5.75,[
     ("심화 가치평가 없음",True,"감정평가식 적정매매가·가치점수 산출 아님"),
     ("설득형 보고서 아님",True,"항목별 등급·근거로 '왜 이 값인가' 설명 X"),
    ],size=11,gap=6)
    pic(s,SHOTS+"bg_brief.png",6.95,1.75,5.95,3.6)
    txt(s,6.95,5.42,5.95,0.3,"▲ 건물개요 브리핑 프린트(상업 템플릿) — 실제 화면",9,False,SUB)
    rect(s,6.95,5.85,5.95,0.85,RGBColor(0xFC,0xF3,0xE2),AMBER)
    txt(s,7.15,5.96,5.6,0.65,[[("정정  ",10.5,True,AMBER),("앞서 '설득 보고서 전무'는 부정확 — 기본 브리핑·수익률 계산은 존재. 부재한 것은 심화 가치평가.",10.5,False,INK)]],10.5)
    foot(s,4)

    # ── 04 · 종합 평가 ──
    s=slide(p); band(s,"부기사",AC1,"04 · 종합 평가")
    closing(s,"부기사",AC1,[
     ("심화 가치평가","브리핑·수익률 계산은 있으나, 감정평가식 적정매매가·가치점수 같은 심화 가치 산출은 다루지 않음."),
     ("설득형 산출물","건물개요 프린트는 있으나, 항목별 등급·근거로 매수자를 설득하는 분석 보고서 형태는 아님."),
     ("진입 비용·확인 한계","상업(부기맨) 초기 400만+연 100만 → 소형 중개인엔 부담. 유료라 내부 동작은 표면 확인."),
    ],"중개 운영·관리(CRM)와 방대한 기능 폭이 강점. 다만 매물의 가격·가치를 심화 평가하고 항목별 근거로 설득하는 분석 산출물은 핵심 영역이 아니다.")
    foot(s,5)
    p.save(OUT+"경쟁사분석_부기사.pptx"); print("saved 부기사 ·",len(p.slides._sldIdLst),"slides")

# ═══════════════════ 2. 디스코 ═══════════════════
def build_disco():
    p=deck(); AC2=GREEN
    cover(p,"디스코","Disco  —  무료 부동산 데이터 조회",
          "지도 기반 실거래가·대장·공시·등기·토지특성 조회 + 경량 매물등록.",AC2,used=True)
    # 실사용 ① 검색·상세
    s=slide(p); band(s,"디스코",AC2,"01 · 직접 써보니"); txt(s,0.5,0.8,10,0.3,"디스코 — 직접 써보니",12,True,AC2); txt(s,0.5,1.13,8,0.4,"건물 하나하나의 원본 열람이 강하다",19,True,INK); usedchip(s,7.05,0.88,True)
    bullets(s,0.5,1.75,6.3,[
     ("검색",True,"'역삼동 601-5' 입력→자동완성 즉시, 지번·도로명·부속지번까지"),
     ("지도",True,"화면 모든 건물에 실거래가·평당가 라벨이 이미 얹혀 있음(압도적 밀도)"),
     ("건물 상세",True,"클릭 한 번→실거래가·경매·토지·건물(대장)·등기 탭 전환"),
     ("실거래 이력",True,"601-5 = 매매 69억('19.02)·40.5억('09.10) 2건 + 등기 링크"),
    ],size=12,gap=7)
    rect(s,0.5,4.95,6.3,1.8,CARD,LINE)
    txt(s,0.72,5.1,5.9,1.55,[
     [("강점  ",11,True,GREEN),("개별 건물의 원본 데이터 깊이 + 서울 전역 커버리지. '동네 시세 감' 잡는 데 탁월.",11,False,INK)],
     [("수익모델  ",11,True,AC),("'우리동네 공인중개사'(PRO 배지) 노출·전화연결 = 중개사 광고/리드",11,False,INK)],
     [("한계  ",11,True,RED),("건물을 '찍어서 여는' 데 최적화 → 조건으로 후보를 좁히는 검색은 약함(다음 장)",11,False,INK)],
    ],11,sp=4)
    pic(s,SHOTS+"disco_detail.png",7.05,1.5,5.9,4.5); txt(s,7.05,6.05,5.9,0.3,"▲ 직접 검색: 역삼동 601-5 상세 — 실거래 69억·거래내역 2건·탭 5종",9.5,False,SUB); foot(s,2)
    # 실사용 ② 조건 검색(필터)의 한계
    s=slide(p); band(s,"디스코",AC2,"02 · 조건 검색"); txt(s,0.5,0.8,10,0.3,"조건 검색은 의외로 얕다",12,True,AC2); txt(s,0.5,1.13,9,0.4,"'찾기'보다 '열람'에 맞춰진 필터",19,True,INK); usedchip(s,0.5,1.7,True)
    txt(s,0.5,2.3,6.3,0.35,[[("필터 패널에서 제공하는 조건 (전부)",12.5,True,INK)]])
    bullets(s,0.5,2.85,6.3,[
     ("부동산 유형",True,"상업용건물·상가·토지·단독 등 10종"),
     ("마커 유형",True,"매물 / 경매 / 실거래가 표시 토글"),
     ("실거래 기간",True,"2006년 ~ 현재"),
     ("거래가격",True,"㎡당 또는 총액 range"),
    ],size=12,gap=7)
    rect(s,0.5,5.35,6.3,1.5,RGBColor(0xFB,0xEE,0xEC),RED)
    txt(s,0.72,5.5,5.9,1.25,[
     [("없는 조건  ",11.5,True,RED),("용도지역 · 대지/연면적 규모 · 건폐율/용적률 · 연식(사용승인) · 도로접면 · 층수 · 주차 · 수익률 · 경사도 등",11,False,INK)],
     [("→ ",11,True,INK),("'유형+가격대'로만 거를 수 있어, 조건으로 매물 후보를 발굴하는 검색엔 부적합.",11,True,INK)],
    ],11,sp=3)
    pic(s,SHOTS+"disco_filter.png",7.05,1.5,5.9,4.55); txt(s,7.05,6.1,5.9,0.3,"▲ 지도 필터 패널 실제 화면 — 유형·마커·기간·가격이 전부",9.5,False,SUB); foot(s,3)
    # 실사용 ② 데이터 깊이
    s=slide(p); band(s,"디스코",AC2,"03 · 데이터 깊이"); txt(s,0.5,0.8,10,0.3,"데이터는 있고, 해석은 없다",12,True,AC2); txt(s,0.5,1.13,9,0.4,"토지특성까지 원본 그대로 노출된다",19,True,INK); usedchip(s,0.5,1.7,True)
    txt(s,0.5,2.25,6.3,0.4,[[("601-5 '토지' 탭에서 실제로 확인한 항목",12.5,True,INK)]])
    bullets(s,0.5,2.85,6.3,[
     ("도로접면",True,"세로한면(가)"),
     ("지형형상",True,"세로장방"),
     ("지형높이",True,"평지"),
     ("용도지역",True,"제3종일반주거지역 / 이용상황 주상기타"),
     ("공시지가·토지이용계획",True,"규제(토지거래허가구역)까지 열람"),
    ],size=12,gap=7)
    rect(s,0.5,5.75,6.3,1.15,RGBColor(0xEF,0xF6,0xEF),GREEN)
    txt(s,0.72,5.9,5.9,0.95,[
     [("핵심 관찰  ",11.5,True,GREEN),("입지·토지특성(도로접면·지형·용도)의 원본 데이터를 빠짐없이 노출한다.",11,False,INK)],
     [("",4,True,INK)],
     [("다만 이를 ",11,True,INK),("하나의 점수·등급이나 추정가로 종합",11,True,RED),("하지는 않는다 — 판단은 사용자 몫.",11,True,INK)],
    ],11,sp=2)
    pic(s,SHOTS+"disco_land.png",7.5,1.5,3.4,5.4); txt(s,7.0,6.95,6.3,0.3,"▲ 601-5 토지 정보 실제 화면(도로접면·지형높이·지형형상)",9.5,False,SUB); foot(s,4)
    # 시사점
    s=slide(p); band(s,"디스코",AC2,"04 · 종합 평가")
    closing(s,"디스코",AC2,[
     ("조건 검색","유형+가격대뿐 — 용도지역·규모·연식·도로·수익률 등으로 후보를 좁히는 발굴형 검색이 약함."),
     ("종합 판단","실거래·공시·토지특성 원본은 풍부하나, 하나의 판단(점수/등급/추정가)으로 묶지 않음."),
     ("관리·산출물","경량 매물등록뿐(CRM 없음) · 화면 조회로 끝 · 매수자에게 건넬 보고서 형태 산출물 없음."),
    ],"개별 건물의 원본 데이터 열람은 강하나, 조건 검색·종합 판단·설득 산출물은 다루지 않는다.")
    foot(s,5)
    p.save(OUT+"경쟁사분석_디스코.pptx"); print("saved 디스코 ·",len(p.slides._sldIdLst),"slides")

# ═══════════════════ 3. 건물닷컴 ═══════════════════
def build_gunmul():
    p=deck(); AC3=AMBER
    cover(p,"건물닷컴","Gunmul.com  —  가격평가 리포트 + 실매물",
          "감정평가 3방식 자동 가격평가 웹리포트 + 컨설팅 + 수익형 실매물.",AC3,used=True)
    # 실사용 — 가격평가 리포트
    s=slide(p); band(s,"건물닷컴",AC3,"01 · 직접 써보니"); txt(s,0.5,0.8,10,0.3,"건물닷컴 — 가격평가 리포트를 열어보니",12,True,AC3); txt(s,0.5,1.13,9,0.4,"감정평가 3방식으로 가격을 삼각측량한다",18.5,True,INK); usedchip(s,0.5,1.68,True)
    rows=[("① 원가기준","토지(공시지가+UP)+건물(복성식)","100,000만"),
          ("② 수익률기준","임대료(지역별 차등)","105,000만"),
          ("③ 실거래기준","거래일+거리+도로+건물가","102,000만")]
    for i,(k,v,val) in enumerate(rows):
        t=2.25+i*0.72; rect(s,0.5,t,6.4,0.62,CARD,LINE); txt(s,0.68,t+0.06,3.0,0.3,k,12,True,AC3)
        txt(s,0.68,t+0.33,4.3,0.25,v,9.5,False,SUB); txt(s,5.1,t+0.16,1.7,0.3,val+"원",11.5,True,INK,PP_ALIGN.RIGHT)
    rect(s,0.5,4.5,6.4,0.62,RGBColor(0xFC,0xF3,0xE2),AMBER)
    txt(s,0.68,4.58,6.1,0.45,[[("최종  ",12,True,AMBER),("3방식 평균 ±7% → 95,170만 ~ 109,496만원 range",11.5,True,INK)]])
    txt(s,0.5,5.35,6.4,1.55,[
     [("느낀 점  ",11,True,GREEN),("감정평가 논리(원가·수익·비교)를 자동화한 건 탄탄. 근거가 투명하게 병렬로 보임.",11,False,INK)],
     [("한계  ",11,True,RED),("의뢰(신청)형 웹페이지 · 홍길동 샘플·2017 카피라이트로 UI 노후 · 항목별 등급·층별 임대·CRM 연계 없음",11,False,INK)],
    ],11,sp=5)
    pic(s,SHOTS+"gunmul_valresult_crop.png",7.1,1.55,5.8,3.0)
    pic(s,SHOTS+"gunmul_appr_preview.png",10.15,4.75,2.7,2.55)
    txt(s,7.1,4.62,3.0,0.3,"▲ 3방식 가격평가결과(실제 미리보기)",9.5,False,SUB)
    txt(s,7.1,5.3,2.9,1.4,[[("좌: 리포트 전체 →  ",9.5,False,SUB)],[("웹페이지 1장짜리 표 형식.",9.5,False,SUB)],[("배포용 산출물 아님.",9.5,True,AMBER)]],9.5); foot(s,2)
    # 시사점
    s=slide(p); band(s,"건물닷컴",AC3,"02 · 시사점")
    closing(s,"건물닷컴",AC3,[
     ("항목별 등급 설명","금액(range)은 주지만 '왜 이 값인가'를 항목별 등급으로 설명하지는 않음."),
     ("워크플로·즉시성","의뢰(신청)형 → 즉석·자기주도 조회가 아님. 중개 CRM과도 분리."),
     ("산출물 형태·UI","웹 1장 표 + 노후 UI. 매수자에게 건넬 배포용 산출물로는 약함."),
    ],"평가 논리는 정통(3방식)이나, 즉시성·항목별 설명·배포용 산출물·중개 워크플로 결합은 다루지 않는다.")
    foot(s,3)
    p.save(OUT+"경쟁사분석_건물닷컴.pptx"); print("saved 건물닷컴 ·",len(p.slides._sldIdLst),"slides")

# ═══════════════════ dispatch ═══════════════════
BUILDERS={"bugisa":build_bugisa,"disco":build_disco,"gunmul":build_gunmul}
if __name__=="__main__":
    sel=[a for a in sys.argv[1:] if a in BUILDERS] or list(BUILDERS)
    for k in sel: BUILDERS[k]()
    print("done ·",len(sel),"deck(s)")
