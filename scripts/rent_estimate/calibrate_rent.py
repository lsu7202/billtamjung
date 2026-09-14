"""임대 추정의 **눈금**을 실측으로 재는 검사(2026-08-28).

물음: 우리 임대 추정이 실제 받는 돈과 자릿수가 맞는가. 순위가 아니라 **금액**을 잰다.

잣대 둘을 교차한다 — 한쪽만으로는 면적 기준을 못 가른다.
  ⓐ 부동산원 상권별 순영업소득(천원/㎡) → 임대수입 원/㎡/월로 환산해 우리 단가와 대조
  ⓑ 실거래가 × 부동산원 소득수익률 → 자산가치 역산으로 ⓐ의 면적 기준을 확정

    backend/.venv/bin/python scripts/rent_estimate/calibrate_rent.py
"""
import asyncio
import os
import statistics as st

import asyncpg
import openpyxl

DSN = os.environ.get("RENT_DSN", "postgresql://postgres:test@localhost:55432/billtamjung")
BASE = os.path.join(os.path.dirname(__file__), "..", "..", "data", "raw")
NOI_SHEETS = {"오피스": "115", "중대형상가": "209", "소규모상가": "309"}
YIELD_SHEETS = {"오피스": "14", "중대형상가": "24", "소규모상가": "34"}
SECT = ['상업용', '업무용', '상업기타', '주상용', '주상기타']


def _stat_file():
    got = [os.path.join(BASE, f) for f in os.listdir(BASE) if "임대동향조사" in f and f.endswith(".xlsx")]
    if not got:
        raise SystemExit(f"임대동향조사 통계표가 없습니다 — {BASE}")
    return max(got, key=os.path.getmtime)


def read_stats():
    """상권별 임대수입(원/㎡/월) + 계열별 서울 소득수익률(연 %) + 운영경비율."""
    wb = openpyxl.load_workbook(_stat_file(), read_only=True)
    rent, expense = {}, {}
    for series, sheet in NOI_SHEETS.items():
        ws = wb[sheet]
        rows = list(ws.iter_rows(min_row=4, values_only=True))
        last = max(i for i, v in enumerate(rows[0]) if isinstance(v, str) and "Q" in str(v))
        acc = {}
        for r in rows[1:]:
            if not r[0] or "서울" not in str(r[0]) or r[last] is None:
                continue
            acc.setdefault(str(r[1]), {})[str(r[2])] = float(r[last])
        for sg, d in acc.items():
            noi, pct, rp = d.get("순영업소득(천원/㎡)"), d.get("순영업소득(%)"), d.get("임대수입(%)")
            if not noi or not pct:
                continue
            rent.setdefault(series, {})[sg] = noi / (pct / 100) * (rp or 100) / 100 / 3 * 1000
            expense.setdefault(series, {})[sg] = 100 - pct
    yld = {}
    for series, sheet in YIELD_SHEETS.items():
        ws = wb[sheet]
        rows = list(ws.iter_rows(min_row=4, values_only=True))
        qs = [i for i, v in enumerate(rows[0]) if isinstance(v, str) and "Q" in str(v)]
        for r in rows[1:]:
            if str(r[0]).strip() == "서울":
                yld[series] = sum(float(r[i]) for i in qs[-4:] if r[i] is not None)
                break
    return rent, expense, yld


async def main():
    rent, expense, yld = read_stats()
    c = await asyncpg.connect(DSN)

    # ⓐ 상권별 단가 배율 — 우리 추정(연면적당) ÷ 실측(임대면적당)
    rows = await c.fetch("""
        SELECT e.sanggwon, e.series, count(*) n,
               percentile_cont(0.5) WITHIN GROUP (ORDER BY e.monthly_rent::float/NULLIF(b.total_area,0)) ours,
               percentile_cont(0.5) WITHIN GROUP (ORDER BY e.rentable_m2::float/NULLIF(b.total_area,0)) share
          FROM master.building_rent_est e JOIN master.buildings b USING(building_pk)
         WHERE e.monthly_rent>0 AND b.total_area>0 AND e.rentable_m2>0 AND e.sanggwon IS NOT NULL
         GROUP BY 1,2 HAVING count(*)>=30""")
    mults, shares = [], []
    print("ⓐ 상권별 단가 배율 (우리 연면적당 ÷ 실측 임대면적당)")
    for r in sorted(rows, key=lambda x: -x["n"])[:20]:
        act = rent.get(r["series"], {}).get(r["sanggwon"])
        if not act:
            continue
        m = r["ours"] / act
        mults.append(m)
        shares.append(r["share"])
        print(f"  {r['sanggwon'][:12]:13s}{r['series'][:5]:6s}{r['n']:6d} "
              f"{r['ours']:9,.0f} / {act:8,.0f} = {m:5.2f}   임대비 {r['share']:.2f}")
    med_m = st.median(mults)
    print(f"  → 배율 중앙 {med_m:.2f} · 임대비(rentable÷연면적) 중앙 {st.median(shares):.2f}\n")

    # ⓑ 면적 기준 확정 — 실거래 평당가 vs 소득수익률 역산 자산가치
    per_m2 = await c.fetchval("""
        SELECT percentile_cont(0.5) WITHIN GROUP (ORDER BY sh.price::float/sh.total_area)
          FROM master.sales_history sh JOIN master.buildings b USING(building_pk)
         WHERE b.bjd_code LIKE '11%' AND sh.contract_ym>='202501' AND sh.price>0
           AND sh.total_area BETWEEN 300 AND 3000 AND b.land_use = ANY($1)""", SECT)
    seoul_noi = rent["중대형상가"]["계"] * 12 * (100 - expense["중대형상가"]["계"]) / 100
    implied = seoul_noi / (yld["중대형상가"] / 100)      # 임대면적당 자산가치
    eff = per_m2 / implied                              # = 임대면적 ÷ 연면적
    print(f"ⓑ 면적 기준 역산 — 실거래 {per_m2:,.0f}원/㎡(연면적) · "
          f"소득수익률 역산 {implied:,.0f}원/㎡\n  → 전용률 {eff:.2f} "
          f"(실측 단가의 분모는 임대면적이고, 연면적의 {eff*100:.0f}%)\n")

    total = med_m / eff
    print(f"=== 총액 배율 {total:.2f}배 — 우리 추정 총임대료가 실측의 {total:.2f}배 ===")
    print(f"    단가 과대 {med_m:.2f}(층별 호가보정) × 전용률 미적용 {1/eff:.2f}\n")

    # 교차검증 — 교정 후 수익률이 실측 소득수익률과 맞는가
    rows = await c.fetch("""
        WITH s AS (SELECT DISTINCT ON (sh.building_pk) sh.price::numeric p,
                          e.annual_rent::numeric ar, e.series
                     FROM master.sales_history sh JOIN master.buildings b USING(building_pk)
                     JOIN master.building_rent_est e USING(building_pk)
                    WHERE b.bjd_code LIKE '11%' AND sh.contract_ym>='202501'
                      AND sh.price>0 AND e.annual_rent>0 AND b.land_use = ANY($1)
                    ORDER BY sh.building_pk, sh.contract_ym DESC)
        SELECT series, count(*) n, percentile_cont(0.5) WITHIN GROUP (ORDER BY ar/p) y
          FROM s GROUP BY 1""", SECT)
    print(f"{'계열':12s}{'n':>5s} {'우리 수익률':>10s} {'교정후':>7s} {'NOI환산':>8s} {'실측':>7s}")
    for r in sorted(rows, key=lambda x: -x["n"]):
        our = float(r["y"]) * 100
        cal = our / total
        noi = cal * (100 - expense[r["series"]]["계"]) / 100
        print(f"{r['series']:12s}{r['n']:5d} {our:9.2f}% {cal:6.2f}% {noi:7.2f}% "
              f"{yld[r['series']]:6.2f}%")
    await c.close()


if __name__ == "__main__":
    asyncio.run(main())
